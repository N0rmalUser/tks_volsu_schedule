from io import BytesIO

from core.enums import WeekType
from openpyxl import load_workbook
from openpyxl.styles import Border, Side

from app.core.constants import TEACHERS_SHEETS_PATH
from app.schemas.schedule import ScheduleEntry


class SpreadsheetsService:
    DAYS_PER_WEEK = 6
    LESSONS_PER_DAY = 7

    START_ROW = 2
    START_COL = 3

    TIMES = {
        1: "08",
        2: "10",
        3: "12",
        4: "13",
        5: "15",
        6: "17",
        7: "18",
    }

    def create_teacher_schedule(
        self,
        lessons: list[ScheduleEntry],
    ) -> BytesIO:
        wb = load_workbook(TEACHERS_SHEETS_PATH)
        ws = wb.active

        not_empty_rows = []

        for lesson in lessons:
            if lesson.lesson_number is None:
                continue

            row = self._get_row(lesson)

            not_empty_rows.extend((row, row + 1) if lesson.week_type == WeekType.ODD else (row, row - 1))

            group = lesson.group or ""

            if lesson.subgroup:
                group += f".{lesson.subgroup}"

            ws.cell(
                row=row,
                column=self.START_COL,
                value=group,
            )

            ws.cell(
                row=row,
                column=self.START_COL + 1,
                value=lesson.room,
            )

            ws.cell(
                row=row,
                column=self.START_COL + 2,
                value=lesson.subject,
            )

        self._delete_empty_rows(ws, not_empty_rows)
        self._merge_equal_cells(ws, self.START_COL + 2)
        self._merge_days(ws)
        self._apply_borders(ws, self.START_COL + 2)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return buffer

    def _apply_borders(self, ws, last_col: int):
        thin = Side(style="thin")
        dotted = Side(style="dotted")
        double = Side(style="double")

        for row in range(self.START_ROW, ws.max_row):
            bottom = dotted if row % 2 == 0 else thin

            for col in range(1, last_col + 1):
                ws.cell(row=row, column=col).border = Border(
                    left=thin,
                    right=thin,
                    bottom=bottom,
                )

            if ws.cell(row=row, column=1).value is not None:
                for col in range(1, last_col + 1):
                    ws.cell(row=row, column=col).border = Border(
                        left=thin,
                        right=thin,
                        top=double,
                        bottom=dotted,
                    )

    def _merge_days(self, ws):
        prev_value = None
        start_merge_row = self.START_ROW

        for row in range(self.START_ROW, ws.max_row + 2):
            value = ws.cell(row=row, column=1).value

            if value != prev_value and prev_value is not None:
                ws.merge_cells(
                    start_row=start_merge_row,
                    start_column=1,
                    end_row=row - 1,
                    end_column=1,
                )
                start_merge_row = row

            prev_value = value

    def _get_row(self, lesson: ScheduleEntry) -> int:
        # day_of_week: 1..6
        # lesson_number: 1..7

        row = self.START_ROW + (lesson.day_of_week - 1) * 14 + (lesson.lesson_number - 1) * 2

        if lesson.week_type != WeekType.ODD:
            row += 1

        return row

    def _delete_empty_rows(self, ws, not_empty_rows: list[int]):
        not_empty_rows = set(not_empty_rows)

        for row in range(ws.max_row, self.START_ROW - 1, -1):
            if row not in not_empty_rows:
                ws.delete_rows(row)

    def _merge_equal_cells(self, ws, last_col: int):
        for row in range(self.START_ROW, ws.max_row, 2):
            for col in range(2, last_col + 1):
                upper = ws.cell(row=row, column=col)
                lower = ws.cell(row=row + 1, column=col)

                if upper.value == lower.value:
                    ws.merge_cells(
                        start_row=row,
                        start_column=col,
                        end_row=row + 1,
                        end_column=col,
                    )

